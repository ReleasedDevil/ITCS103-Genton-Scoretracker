import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl import Workbook
import os

EXCEL_FILE = "student_scores.xlsx"

def create_excel_file():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Score", "Status"])
        wb.save(EXCEL_FILE)

def get_status(score):
    return "Pass" if score >= 75 else "Fail"

def save_record():
    name = name_entry.get().strip()
    try:
        score = int(score_entry.get())
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a number.")
        return

    if not name:
        messagebox.showerror("Invalid Input", "Name cannot be empty.")
        return

    status = get_status(score)

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append([name, score, status])
    wb.save(EXCEL_FILE)
    messagebox.showinfo("Success", f"Record saved for {name}.")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

def show_records():
    if not os.path.exists(EXCEL_FILE):
        messagebox.showerror("Error", "No data found.")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

    records = "\n".join(
        f"{row[0].value}: {row[1].value} - {row[2].value}" for row in ws.iter_rows(min_row=2)
    )

    if not records:
        messagebox.showinfo("Records", "No records found.")
    else:
        messagebox.showinfo("All Student Records", records)


create_excel_file()

window = tk.Tk()
window.title("Student Score Tracker")
window.geometry("400x320")
window.configure(bg="#e6f2ff") 

label_font = ("Segoe UI", 11)
entry_font = ("Segoe UI", 11)
button_font = ("Segoe UI", 10, "bold")

frame = tk.Frame(window, bg="#e6f2ff")
frame.pack(pady=30)

tk.Label(frame, text="Student Name:", font=label_font, bg="#e6f2ff").grid(row=0, column=0, sticky="w", pady=(0,5))
name_entry = tk.Entry(frame, font=entry_font, width=30, bg="white", relief="solid", bd=1)
name_entry.grid(row=1, column=0, pady=(0,15))

tk.Label(frame, text="Score:", font=label_font, bg="#e6f2ff").grid(row=2, column=0, sticky="w", pady=(0,5))
score_entry = tk.Entry(frame, font=entry_font, width=30, bg="white", relief="solid", bd=1)
score_entry.grid(row=3, column=0, pady=(0,15))

save_btn = tk.Button(frame, text="Save Record", font=button_font, bg="#4CAF50", fg="white", width=25, command=save_record)
save_btn.grid(row=4, column=0, pady=10)

show_btn = tk.Button(frame, text="Show All Records", font=button_font, bg="#2196F3", fg="white", width=25, command=show_records)
show_btn.grid(row=5, column=0)

window.mainloop()